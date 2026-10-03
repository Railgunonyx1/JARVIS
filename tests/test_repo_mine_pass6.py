"""Tests for repo-mine pass-6 tools: data (pandas), docx reports, routines.

Hermetic: CSV/JSON fixtures are written into tmp_path, ProjectContext is
monkeypatched to that root in each tool module, and routine.run is asserted
to return a *plan* rather than executing anything.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

import pandas as pd
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from tools import build_default_registry  # noqa: E402

# ----------------------------------------------------------------- registry


def test_pass6_tools_registered():
    registry = build_default_registry()
    for name in (
        "data.stats", "data.query", "data.convert",
        "doc.report",
        "routine.add", "routine.list", "routine.remove", "routine.run",
    ):
        tool = registry.get(name)
        assert tool is not None, f"{name} missing from registry"
        assert tool.handler is not None


# ------------------------------------------------------------------ fixtures


@pytest.fixture()
def data_root(tmp_path, monkeypatch):
    """Project root with two data files; all three tool modules point at it."""
    pd.DataFrame(
        {"region": ["EMEA", "EMEA", "APAC", "AMER"], "revenue": [100, 2500, 700, 1200],
         "units": [10, 22, 7, 12]}
    ).to_csv(tmp_path / "sales.csv", index=False)
    (tmp_path / "people.jsonl").write_text(
        '\n'.join(json.dumps(r) for r in
                  [{"name": "Ada", "age": 36}, {"name": "Ben", "age": 41}]),
        encoding="utf-8",
    )
    fake = type("C", (), {"root_path": tmp_path})
    from core.project import ProjectContext

    monkeypatch.setattr(
        ProjectContext, "discover",
        staticmethod(lambda cwd=None: fake()),
    )
    return tmp_path


# --------------------------------------------------------------------- data


def test_data_stats_profiles_columns(data_root):
    from tools.data_tools import data_stats

    res = data_stats({"path": "sales.csv"})
    assert res.success
    assert "rows: 4" in res.output
    assert "region" in res.output and "revenue" in res.output
    assert res.metadata["rows"] == 4


def test_data_stats_correlation_detected(data_root):
    from tools.data_tools import data_stats

    res = data_stats({"path": "sales.csv"})
    assert "strongest numeric correlation" in res.output


def test_data_stats_missing_file(data_root):
    from tools.data_tools import data_stats

    res = data_stats({"path": "nope.csv"})
    assert not res.success
    assert "Not found" in res.error


def test_data_query_filters_rows(data_root):
    from tools.data_tools import data_query

    res = data_query({"path": "sales.csv", "expr": "revenue > 1000"})
    assert res.success
    assert "2 of 4 rows match" in res.output
    assert res.metadata["matched"] == 2


def test_data_query_bad_expr_reports_columns(data_root):
    from tools.data_tools import data_query

    res = data_query({"path": "sales.csv", "expr": "bogus_col > 1"})
    assert not res.success
    assert "region" in res.error  # column list included for the model


def test_data_convert_csv_to_jsonl(data_root):
    from tools.data_tools import data_convert

    res = data_convert({"path": "sales.csv", "dest": "out/sales.jsonl"})
    assert res.success, res.error
    out = data_root / "out" / "sales.jsonl"
    lines = out.read_text(encoding="utf-8").strip().splitlines()
    assert len(lines) == 4
    assert json.loads(lines[0])["region"] == "EMEA"


def test_data_convert_rejects_outside_root(data_root):
    from tools.data_tools import data_convert

    res = data_convert({"path": "sales.csv", "dest": "C:/Windows/temp/x.csv"})
    assert not res.success
    assert "project root" in res.error


def test_data_tools_unsupported_type(data_root):
    from tools.data_tools import data_query

    (data_root / "x.parquet").write_bytes(b"")
    res = data_query({"path": "x.parquet", "expr": "a > 1"})
    assert not res.success


# ---------------------------------------------------------------- doc.report


def test_doc_report_builds_document(data_root):
    from docx import Document

    from tools.doc_report import doc_report

    res = doc_report({
        "title": "Quarterly Review",
        "subtitle": "Prepared by JARVIS",
        "sections": [
            {"title": "Summary", "text": "Revenue grew.\n\nCosts held flat."},
            {"title": "Risks", "bullets": ["FX exposure", "Churn in EMEA"]},
            "Appendix",
        ],
        "dest": "reports/q2.docx",
    })
    assert res.success, res.error
    out = data_root / "reports" / "q2.docx"
    assert out.exists()

    doc = Document(str(out))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert "Summary" in headings and "Risks" in headings and "Appendix" in headings
    titles = [p.text for p in doc.paragraphs if p.style.name == "Title"]
    assert "Quarterly Review" in titles
    bullets = [p.text for p in doc.paragraphs if p.style.name == "List Bullet"]
    assert "FX exposure" in bullets and "Churn in EMEA" in bullets


def test_doc_report_requires_title_and_sections(data_root):
    from tools.doc_report import doc_report

    assert not doc_report({"title": "", "sections": [{"title": "x"}]}).success
    assert not doc_report({"title": "T", "sections": []}).success


def test_doc_report_outside_root_rejected(data_root):
    from tools.doc_report import doc_report

    res = doc_report({"title": "T", "sections": [{"title": "s"}], "dest": "C:/Windows/temp/r.docx"})
    assert not res.success
    assert "project root" in res.error


# ------------------------------------------------------------------ routines


def _add(rt, name, steps):
    return rt.routine_add({"name": name, "steps": steps})


def test_routine_add_list_remove_cycle(data_root):
    from tools import routines as rt

    res = _add(rt, "digest", [
        {"tool": "web.search", "args": {"query": "AI news"}},
        {"tool": "notify.send", "args": {"message": "done"}},
    ])
    assert res.success, res.error

    listed = rt.routine_list({})
    assert "digest" in listed.output and "web.search" in listed.output

    assert rt.routine_remove({"name": "digest"}).success
    assert "No routines stored" in rt.routine_list({}).output


def test_routine_add_validates_steps(data_root):
    from tools import routines as rt

    assert not rt.routine_add({"name": "x", "steps": []}).success
    assert not rt.routine_add({"name": "x", "steps": [{"args": {}}]}).success
    assert not rt.routine_add({"name": "", "steps": [{"tool": "web.search"}]}).success


def test_routine_run_returns_validated_plan_not_execution(data_root, monkeypatch):
    """routine.run must produce a plan, never fire the referenced tools."""
    from tools import routines as rt

    _add(rt, "safe", [{"tool": "web.search", "args": {"query": "q"}}])

    fired = []
    import tools.web_search as ws
    monkeypatch.setattr(ws, "web_search", lambda a: fired.append(a) or ws.web_search(a))

    res = rt.routine_run({"name": "safe"})
    assert res.success, res.error
    assert fired == [], "routine.run executed a tool directly — boundary bypass"
    plan = res.metadata["plan"]
    assert plan[0]["tool"] == "web.search"
    assert "Execute each step in order" in res.output


def test_routine_run_counts_runs(data_root):
    from tools import routines as rt

    _add(rt, "counter", [{"tool": "system.status", "args": {}}])
    rt.routine_run({"name": "counter"})
    rt.routine_run({"name": "counter"})
    data = json.loads((data_root / "memory" / "routines.json").read_text(encoding="utf-8"))
    assert data["counter"]["runs"] == 2


def test_routine_run_unknown_tool_fails_validation(data_root):
    from tools import routines as rt

    _add(rt, "broken", [{"tool": "not_a_real_tool", "args": {}}])
    res = rt.routine_run({"name": "broken"})
    assert not res.success
    assert "unknown tools" in res.error.lower()
    assert "not_a_real_tool" in res.error


def test_routine_run_unknown_name_lists_known(data_root):
    from tools import routines as rt

    _add(rt, "known", [{"tool": "system.status", "args": {}}])
    res = rt.routine_run({"name": "missing"})
    assert not res.success
    assert "known" in res.error
