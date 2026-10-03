"""Report generation — Word documents via python-docx.

Derived from the repo-mine list (python-openxml/python-docx): research
results and task summaries can land as a formatted .docx instead of raw
markdown the user has to convert themselves. Writes only inside the
project root.
"""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Any

from tools.schema import ToolResult

logger = logging.getLogger("jarvis.tools.doc_report")

MAX_SECTIONS = 40
MAX_ITEM_CHARS = 4000


def _root() -> Path:
    from core.project import ProjectContext
    return ProjectContext.discover().root_path


def _resolve(path_str: str) -> Path:
    p = Path(path_str).expanduser()
    if not p.is_absolute():
        p = _root() / p
    return p


def _normalise_sections(raw: Any) -> list[dict[str, Any]]:
    """Accept sections as list-of-dicts or a list of strings (title only)."""
    if isinstance(raw, str):
        return [{"title": raw}]
    out: list[dict[str, Any]] = []
    for item in list(raw or [])[:MAX_SECTIONS]:
        if isinstance(item, str):
            out.append({"title": item})
        elif isinstance(item, dict):
            out.append(item)
    return out


def doc_report(args: dict[str, Any]) -> ToolResult:
    """Build a .docx report: title page, headings, paragraphs and bullet lists."""
    try:
        from docx import Document
        from docx.shared import Pt
    except ImportError:
        return ToolResult(success=False, error="python-docx not installed (pip install python-docx)")

    title = str(args.get("title", "")).strip()
    if not title:
        return ToolResult(success=False, error="title is required")
    dest_arg = str(args.get("dest", "")).strip()
    if not dest_arg:
        dest_arg = "report.docx"
    dest = _resolve(dest_arg)
    if dest.suffix.lower() != ".docx":
        dest = dest.with_suffix(".docx")
    try:
        dest.resolve().relative_to(_root().resolve())
    except ValueError:
        return ToolResult(success=False, error=f"Destination must be inside the project root: {dest}")

    sections = _normalise_sections(args.get("sections"))
    if not sections:
        return ToolResult(
            success=False,
            error=(
                "sections is required: a list of {'title', 'text'?, 'bullets'?} "
                "dicts (or plain strings for heading-only sections)"
            ),
        )

    try:
        doc = Document()
        doc.add_heading(title, level=0)

        author = str(args.get("author", "")).strip()
        subtitle = str(args.get("subtitle", "")).strip()
        if subtitle or author:
            p = doc.add_paragraph()
            run = p.add_run(subtitle or author)
            run.italic = True

        for sec in sections:
            sec_title = str(sec.get("title", "")).strip()
            if sec_title:
                doc.add_heading(sec_title, level=1)
            text = sec.get("text")
            if text:
                for chunk in str(text).split("\n\n"):
                    chunk = chunk.strip()
                    if chunk:
                        doc.add_paragraph(chunk)
            bullets = sec.get("bullets") or []
            if isinstance(bullets, str):
                bullets = [b for b in (ln.strip() for ln in bullets.split("\n")) if b]
            for b in list(bullets)[:100]:
                b = str(b)[:MAX_ITEM_CHARS]
                doc.add_paragraph(b, style="List Bullet")

        dest.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(dest))
    except Exception as e:  # noqa: BLE001
        return ToolResult(success=False, error=f"Report build failed: {e}")

    rel = dest.relative_to(_root())
    return ToolResult(
        success=True,
        output=f"Report written: {rel} ({len(sections)} sections)",
        metadata={"path": str(rel), "sections": len(sections)},
    )


__all__ = ["doc_report"]
